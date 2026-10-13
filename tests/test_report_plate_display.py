from core import reports
import re
import zlib
from pathlib import Path


def _stub_report_database(monkeypatch):
    monkeypatch.setattr(reports.db, "get_case_policy_records", lambda _vid: [])
    monkeypatch.setattr(reports.db, "is_case_confirmed", lambda _vid: True)
    monkeypatch.setattr(reports.db, "is_notice_printed", lambda _vid: False)


def test_report_includes_verified_plate_text(monkeypatch):
    _stub_report_database(monkeypatch)

    row = reports._enrich_report_row(
        {
            "id": 12,
            "violation_type": "Illegal Parking",
            "plate_status": "recognized",
            "plate_text": "JAD7419",
        }
    )

    assert row["plate_display"] == "JAD7419 (verified)"


def test_report_labels_machine_text_as_unverified_candidate(monkeypatch):
    _stub_report_database(monkeypatch)
    monkeypatch.setattr(
        "core.plate_review.machine_result_for_violation",
        lambda _db, _vid: {
            "primary_candidate_id": "candidate-1",
            "candidates": [{"candidate_id": "candidate-1", "ocr_raw": "JAD7419"}],
            "association_uncertain": False,
        },
    )

    row = reports._enrich_report_row(
        {
            "id": 12,
            "violation_type": "Illegal Parking",
            "plate_status": "not_attempted",
            "plate_text": None,
        }
    )

    assert row["plate_display"] == "JAD7419 (OCR candidate; unverified)"
    assert "plate_display" in [key for key, _header, _width in reports._COLUMNS]
    assert dict((key, header) for key, header, _width in reports._COLUMNS)["legal_status"] == (
        "Legal Mapping Status"
    )


def test_report_marks_uncertain_candidate_association(monkeypatch):
    _stub_report_database(monkeypatch)
    monkeypatch.setattr(
        "core.plate_review.machine_result_for_violation",
        lambda _db, _vid: {
            "primary_candidate_id": "candidate-1",
            "candidates": [{"candidate_id": "candidate-1", "ocr_raw": "JAD7419"}],
            "association_uncertain": True,
        },
    )

    row = reports._enrich_report_row(
        {"id": 12, "violation_type": "Illegal Parking", "plate_status": "not_attempted"}
    )

    assert row["plate_display"] == "JAD7419 (OCR candidate; unverified; association uncertain)"


def test_report_separates_detector_and_rule_scores_and_keeps_zero(monkeypatch):
    _stub_report_database(monkeypatch)
    row = reports._enrich_report_row({
        "id": 12,
        "violation_type": "Counterflow",
        "confidence": 0.0,
        "violation_confidence": 0.0,
        "detection_confidence": 0.0,
    })
    columns = dict((key, header) for key, header, _width in reports._COLUMNS)
    assert columns["detection_confidence"] == "Detector Model Score"
    assert columns["violation_confidence"] == "Rule Evidence Score"
    assert reports._cell_value(row, "detection_confidence") == "0%"
    assert reports._cell_value(row, "violation_confidence") == "0%"
    assert reports._cell_value({"violation_confidence": None, "confidence": None}, "violation_confidence") == "-"


def test_report_keeps_category_and_legal_status_paired_per_contributing_behavior(monkeypatch):
    _stub_report_database(monkeypatch)
    monkeypatch.setattr(
        reports.db,
        "get_case_policy_records",
        lambda _vid: [
            {
                "canonical_rule": "Illegal Parking",
                "official_category": "Parking Offense",
                "legal_status": "unverified",
            },
            {
                "canonical_rule": "Obstruction",
                "official_category": "Obstruction of Traffic Flow",
                "legal_status": "verified",
            },
        ],
    )

    row = reports._enrich_report_row({"id": 12, "violation_type": "Illegal Parking"})

    assert row["official_category"] == (
        "Illegal Parking: Parking Offense [unverified]; "
        "Obstruction: Obstruction of Traffic Flow [verified]"
    )
    assert row["legal_status"] == (
        "Illegal Parking: unverified; Obstruction: verified"
    )


def test_excel_export_contains_plate_column_and_candidate_label(tmp_path):
    row = {
        "id": 12,
        "violation_type": "Illegal Parking",
        "official_category": "Obstruction of Traffic Flow",
        "legal_status": "unverified",
        "plate_display": "JAD7419 (OCR candidate; unverified)",
        "case_outcome": "case_confirmed",
        "vehicle_class": "Passenger Vehicle",
        "confidence": 0.87,
        "detected_at": "2026-10-04 07:26:10",
        "status": "confirmed",
    }
    output = tmp_path / "report.xlsx"

    reports._generate_excel([row], "test", {}, output)

    from openpyxl import load_workbook

    worksheet = load_workbook(output, read_only=True)["Violations"]
    headers = [cell.value for cell in worksheet[5]]
    plate_column = headers.index("Plate (verification)") + 1
    assert worksheet.cell(row=6, column=plate_column).value == row["plate_display"]


def test_pdf_export_generates_with_long_unverified_plate_label(tmp_path):
    row = {
        "id": 12,
        "violation_type": "Illegal Parking",
        "official_category": "Obstruction of Traffic Flow",
        "legal_status": "unverified",
        "plate_display": "JAD7419 (OCR candidate; unverified; association uncertain)",
        "case_outcome": "case_confirmed",
        "vehicle_class": "Passenger Vehicle",
        "confidence": 0.87,
        "detected_at": "2026-10-04 07:26:10",
        "status": "confirmed",
    }
    output = tmp_path / "report.pdf"

    reports._generate_pdf([row], "test", {}, output)

    assert output.is_file()
    assert output.stat().st_size > 0


def test_pdf_keeps_fused_rows_together_and_splits_oversized_rows_without_losing_text(tmp_path):
    rows = [
        {
            "id": 100 + index,
            "violation_type": "Illegal Parking + Obstruction",
            "official_category": (
                "Illegal Parking: Parking Offense [unverified]; "
                "Obstruction: Obstruction of Traffic Flow [verified]"
            ),
            "legal_status": "Illegal Parking: unverified; Obstruction: verified",
            "plate_display": "JAD7419 (OCR candidate; unverified; association uncertain)",
            "case_outcome": "case_confirmed",
            "vehicle_class": "Passenger Vehicle",
            "confidence": 0.87,
            "detected_at": "2026-10-04 07:26:10",
            "status": "confirmed",
        }
        for index in range(12)
    ]
    rows.append({
        **rows[0],
        "id": 999,
        "official_category": "Evidence phrase " * 180,
        "legal_status": "Detailed status phrase " * 120,
        "plate_display": "Long verification detail " * 100,
    })
    output = tmp_path / "fused-pagination.pdf"

    reports._generate_pdf(rows, "pagination", {}, output)

    data = output.read_bytes()
    streams = re.findall(rb"stream\r?\n(.*?)\r?\nendstream", data, re.DOTALL)
    rendered_text = b"\n".join(
        zlib.decompress(stream) for stream in streams
        if stream and stream.startswith(b"x")
    )
    page_count = data.count(b"/Type /Page\n") + data.count(b"/Type /Page ")
    assert data.startswith(b"%PDF")
    assert page_count < 20
    for case_id in [str(100 + index) for index in range(12)] + ["999"]:
        assert f"({case_id}) Tj".encode("ascii") in rendered_text
    assert b"Parking Offense" in rendered_text
    assert b"association" in rendered_text
    assert b"uncertain" in rendered_text
    assert b"Long" in rendered_text
    assert b"verification" in rendered_text
    assert b"detail" in rendered_text


def test_pdf_export_handles_empty_report(tmp_path):
    output = tmp_path / "empty-report.pdf"

    reports._generate_pdf([], "empty", {}, output)

    assert output.is_file()
    assert output.read_bytes().startswith(b"%PDF")


def test_pdf_and_excel_embed_the_login_brand_mark(tmp_path):
    from openpyxl import load_workbook

    row = {"id": 1, "violation_type": "Illegal Parking", "status": "confirmed"}
    pdf_path = tmp_path / "branded-report.pdf"
    xlsx_path = tmp_path / "branded-report.xlsx"

    reports._generate_pdf([row], "test", {}, pdf_path)
    reports._generate_excel([row], "test", {}, xlsx_path)

    pdf_bytes = pdf_path.read_bytes()
    workbook = load_workbook(xlsx_path)
    assert b"/Subtype /Image" in pdf_bytes
    assert len(workbook["Violations"]._images) == 1
    assert workbook["Violations"]["B1"].value.startswith("TAVIDM Traffic Violation Report")
    sheet = workbook["Violations"]
    image = sheet._images[0]
    assert image.width > 0 and image.height > 0
    assert image.width / image.height == 1
    assert sheet.page_setup.orientation == "landscape"
    assert sheet.page_setup.fitToWidth == 1
    assert sheet.page_setup.fitToHeight == 0
    assert sheet.print_title_rows == "$1:$5"
    assert sheet.print_area == "'Violations'!$A$1:$K$6"


def test_report_logo_missing_fails_before_output_or_report_record(tmp_path, monkeypatch):
    output_dir = tmp_path / "reports"
    monkeypatch.setattr(reports, "REPORTS_FOLDER", output_dir)
    monkeypatch.setattr(reports, "_REPORT_LOGO", tmp_path / "missing.png")
    monkeypatch.setattr(reports, "_fetch_rows", lambda _filters: [])
    inserted = []
    monkeypatch.setattr(reports.db, "insert_report", lambda **kwargs: inserted.append(kwargs))

    import pytest

    with pytest.raises(reports.ReportLogoError, match="approved TAVIDM report logo"):
        reports.generate_report("pdf")

    assert not output_dir.exists()
    assert inserted == []


def test_report_render_failure_removes_partial_output_without_inserting_record(tmp_path, monkeypatch):
    output_dir = tmp_path / "reports"
    monkeypatch.setattr(reports, "REPORTS_FOLDER", output_dir)
    monkeypatch.setattr(reports, "_fetch_rows", lambda _filters: [])
    inserted = []
    monkeypatch.setattr(reports.db, "insert_report", lambda **kwargs: inserted.append(kwargs))

    def write_then_fail(_rows, _title, _filters, out_path):
        out_path.parent.mkdir(parents=True, exist_ok=True)
        out_path.write_bytes(b"partial")
        raise RuntimeError("render failed")

    monkeypatch.setattr(reports, "_generate_pdf", write_then_fail)
    import pytest

    with pytest.raises(RuntimeError, match="render failed"):
        reports.generate_report("pdf")

    assert list(output_dir.glob("*")) == []
    assert inserted == []


def test_report_lookup_failure_keeps_the_inserted_record_file_pair(tmp_path, monkeypatch):
    output_dir = tmp_path / "reports"
    monkeypatch.setattr(reports, "REPORTS_FOLDER", output_dir)
    monkeypatch.setattr(reports, "_fetch_rows", lambda _filters: [])
    inserted = []

    def insert_report(**kwargs):
        inserted.append(kwargs)
        return 44

    monkeypatch.setattr(reports.db, "insert_report", insert_report)
    monkeypatch.setattr(
        reports.db, "get_report", lambda _report_id: (_ for _ in ()).throw(RuntimeError("lookup failed"))
    )

    import pytest

    with pytest.raises(RuntimeError, match="lookup failed"):
        reports.generate_report("pdf")

    assert len(inserted) == 1
    assert (Path(inserted[0]["file_path"])).is_file()


def test_corrupt_report_logo_fails_as_controlled_error_without_partial_file(tmp_path, monkeypatch):
    corrupt = tmp_path / "corrupt.png"
    corrupt.write_bytes(b"not an image")
    monkeypatch.setattr(reports, "_REPORT_LOGO", corrupt)
    output = tmp_path / "report.pdf"

    import pytest

    with pytest.raises(reports.ReportLogoError, match="approved TAVIDM report logo"):
        reports._generate_pdf([], "empty", {}, output)

    assert not output.exists()


def test_report_api_returns_actionable_service_error_for_missing_brand(enforcer_client, monkeypatch):
    import app as flask_app

    def fail_for_logo(**_kwargs):
        raise reports.ReportLogoError(
            "The approved TAVIDM report logo is missing or unreadable. Restore "
            "static/images/tavidm-report-logo.png, then retry."
        )

    monkeypatch.setattr(flask_app.reports_core, "generate_report", fail_for_logo)

    response = enforcer_client.post("/api/reports", json={"format": "pdf"})

    assert response.status_code == 503
    assert "static/images/tavidm-report-logo.png" in response.get_json()["error"]


def test_pdf_header_labels_wrap_inside_landscape_column_geometry():
    from fpdf import FPDF

    pdf = FPDF(orientation="L")
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 7.5)
    widths = [9, 33, 45, 24, 35, 20, 22, 14, 14, 27, 17]
    headers = [header for _key, header, _width in reports._COLUMNS]
    assert sum(widths) <= pdf.w - pdf.l_margin - pdf.r_margin
    wrapped = [
        reports._wrap_pdf_cell_text(pdf, header, width + 2)
        for header, width in zip(headers, widths)
    ]
    for lines, width in zip(wrapped, widths):
        usable = width - 2 * pdf.c_margin
        assert all(pdf.get_string_width(line) <= usable for line in lines)
    header_height = max(8.0, max(len(lines) for lines in wrapped) * 4.0 + 2.0)
    assert header_height >= max(len(lines) for lines in wrapped) * 4.0
