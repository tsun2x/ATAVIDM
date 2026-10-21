from pathlib import Path
import ast, hashlib, json, re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(r'D:\tavidm')
OUT = Path(__file__).parent
im=Image.new('RGB',(1800,560),'white');draw=ImageDraw.Draw(im)
font=ImageFont.truetype(r'C:\Windows\Fonts\arial.ttf',28)
bold=ImageFont.truetype(r'C:\Windows\Fonts\arialbd.ttf',32)
for x,title,body in [(15,'INPUT',['Authorized video or camera','Scene geometry and time','Rules and model capabilities']), (625,'PROCESS',['Detect and track objects','Evaluate rules and save evidence','Review and record decisions']), (1235,'OUTPUT',['Candidates and reviewed records','Replay and passage counts','Analytics and reports'])]:
    draw.rounded_rectangle((x,35,x+550,365),radius=12,fill='#F1F4F7',outline='#3F4C59',width=3)
    draw.text((x+25,65),title,font=bold,fill='black')
    for j,t in enumerate(body):draw.text((x+25,140+j*65),t,font=font,fill='black')
for x in [570,1180]:
    draw.line((x,200,x+48,200),fill='black',width=4)
    draw.polygon([(x+48,200),(x+34,191),(x+34,209)],fill='black')
draw.line([(1500,380),(1500,440),(285,440),(285,380)],fill='#3F4C59',width=3)
draw.polygon([(285,380),(277,396),(293,396)],fill='#3F4C59')
draw.text((330,474),'Authorized feedback to annotations and settings for a new run',font=font,fill='black')
im.save(OUT/'conceptual_framework.png')
inventory = []
paths = [ROOT/'app.py', ROOT/'config.py']
for folder, pattern in [('core','*.py'),('database','*.py'),('templates','*.html'),('static/js','*.js'),('tests','test_*.py'),('tests/js','*.cjs'),('scripts','*.py'),('database','*.sql')]:
    paths.extend((ROOT/folder).rglob(pattern))
paths.extend([ROOT/'config/training/class_schema.json', ROOT/'requirements.txt'])
for path in sorted(set(paths)):
    raw=path.read_bytes()
    item={'path':path.relative_to(ROOT).as_posix(),'sha256':hashlib.sha256(raw).hexdigest(),'lines':len(raw.splitlines())}
    if path.suffix=='.py':
        try:
            tree=ast.parse(raw.decode('utf-8-sig'))
            item['purpose']=ast.get_docstring(tree)
            item['definitions']=[{'name':n.name,'line':n.lineno} for n in ast.walk(tree) if isinstance(n,(ast.FunctionDef,ast.ClassDef,ast.AsyncFunctionDef))]
        except Exception as e: item['parse_error']=str(e)
    inventory.append(item)
(OUT/'source_inventory.json').write_text(json.dumps({'date':'2026-10-11','method':'Static inventory and AST inspection; no application imports or runtime validation','files':inventory},indent=2),encoding='utf-8')

doc=Document()
for style in doc.styles:
    for border in style.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
s=doc.sections[0]
s.page_width=Inches(8.5); s.page_height=Inches(11)
s.top_margin=s.bottom_margin=Inches(.8)
s.left_margin=s.right_margin=Inches(.9)
s.header_distance=s.footer_distance=Inches(.35)
for name in ['Normal','Body Text','Title','Subtitle','Heading 1','Heading 2','Heading 3','List Bullet','List Number']:
    st=doc.styles[name]; st.font.name='Times New Roman'; st.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(11.5)
doc.styles['Normal'].paragraph_format.line_spacing=1.13
doc.styles['Normal'].paragraph_format.space_after=Pt(7)
for name,size in [('Title',23),('Heading 1',16),('Heading 2',12.5),('Heading 3',11.5)]:
    doc.styles[name].font.size=Pt(size)
    doc.styles[name].font.bold=True
    doc.styles[name].paragraph_format.space_before=Pt(10)
    doc.styles[name].paragraph_format.space_after=Pt(7)
    doc.styles[name].paragraph_format.keep_with_next=True
hp=s.header.paragraphs[0]; hp.text='TAVIDM  |  Chapter 3 and system procedures'
hp.runs[0].font.size=Pt(9)
fp=s.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
fp.add_run('Source based draft  •  11 October 2026  |  ')
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); fp._p.append(fld)
for r in fp.runs:r.font.size=Pt(9)
doc.core_properties.title='TAVIDM Chapter 3 and System SOP'
doc.core_properties.subject='Technical Background and Standard Operating Procedure'
doc.core_properties.author='TAVIDM Project Team'
doc.core_properties.keywords='TAVIDM, Chapter 3, methodology, SOP, source based draft'

def add_table(rows):
    table=doc.add_table(rows=1,cols=len(rows[0])); table.alignment=WD_TABLE_ALIGNMENT.CENTER
    table.autofit=False
    widths=([2.0,4.7] if len(rows[0])==2 else [1.7,1.1,3.9] if len(rows[0])==3 else [1.3]*(len(rows[0])))
    if rows[0][0]=='ID':widths=[.5,6.2]
    for c,w in zip(table.columns,widths):c.width=Inches(w)
    for i,row in enumerate(rows):
        cells=table.rows[0].cells if i==0 else table.add_row().cells
        for j,(cell,value) in enumerate(zip(cells,row)):
            cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text=value.strip()
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4);p.paragraph_format.line_spacing=1.05
                for r in p.runs:r.font.size=Pt(10);r.bold=(i==0)
            pr=cell._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            pr.append(borders)
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');margins.append(e)
            pr.append(margins)
            if i==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EDF2');pr.append(shade)
        trpr=table.rows[i]._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if i==0:trpr.append(OxmlElement('w:tblHeader'))
    doc.add_paragraph().paragraph_format.space_after=Pt(1)

text=(OUT/'manuscript.md').read_text(encoding='utf-8')
lines=text.splitlines();i=0
while i<len(lines):
    line=lines[i].strip()
    if not line:i+=1;continue
    if line=='<!-- pagebreak -->':
        upcoming=next((x.strip() for x in lines[i+1:] if x.strip()),'')
        if upcoming.startswith(('## CHAPTER','## Appendix A','## Appendix B')):doc.add_page_break()
    elif line=='[FIGURE]':
        doc.add_picture(str(OUT/'conceptual_framework.png'),width=Inches(6.7))
    elif line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].strip().startswith('|'):
            row=[v.strip() for v in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r'[-: ]+',v) for v in row):rows.append(row)
            i+=1
        add_table(rows);continue
    elif line.startswith('# '):doc.add_paragraph(line[2:],'Title')
    elif line.startswith('## '):doc.add_paragraph(line[3:],'Heading 1')
    elif line.startswith('### '):doc.add_paragraph(line[4:],'Heading 2')
    elif line.startswith('- '):doc.add_paragraph(line[2:],'List Bullet')
    else:
        p=doc.add_paragraph(line)
        p.paragraph_format.widow_control=True
    i+=1
file=OUT/'TAVIDM_Chapter_3_and_SOP.docx';doc.save(file)
print(json.dumps({'document':str(file),'inventory_files':len(inventory),'words':len(text.split()),'tables':len(doc.tables),'paragraphs':len(doc.paragraphs)}))
